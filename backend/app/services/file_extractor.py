import io
import os
import re
from typing import Optional, Dict, Any
from app.core.config import settings


class FileValidationError(Exception):
    """Raised when uploaded file fails security validation."""
    pass


class ExtractedDocument:
    def __init__(
        self,
        filename: str,
        file_type: str,
        size_bytes: int,
        text: str,
        page_count: int,
        word_count: int,
        character_count: int,
        excerpt: str,
        metadata: Optional[Dict[str, Any]] = None
    ):
        self.filename = filename
        self.file_type = file_type
        self.size_bytes = size_bytes
        self.text = text
        self.page_count = page_count
        self.word_count = word_count
        self.character_count = character_count
        self.excerpt = excerpt
        self.metadata = metadata or {}


class FileExtractor:
    """
    Safely validates file format, enforces strict size limits, rejects executable payloads,
    and extracts plain text in-memory from TXT, PDF, and DOCX files.
    """

    ALLOWED_EXTENSIONS = {".txt", ".pdf", ".docx"}
    DANGEROUS_MAGIC_HEADERS = [
        (b"MZ", "Windows PE Executable"),
        (b"\x7fELF", "Linux ELF Executable"),
        (b"\xca\xfe\xba\xbe", "Java Bytecode / Mach-O"),
        (b"#!", "Shell Script"),
    ]

    @classmethod
    def validate_and_extract(cls, file_bytes: bytes, filename: str) -> ExtractedDocument:
        # 1. Validate file size
        max_bytes = settings.max_file_size_mb * 1024 * 1024
        if len(file_bytes) > max_bytes:
            raise FileValidationError(
                f"File size ({len(file_bytes) / 1024 / 1024:.2f} MB) exceeds maximum allowed limit of {settings.max_file_size_mb} MB"
            )

        if len(file_bytes) == 0:
            raise FileValidationError("Uploaded file is empty (0 bytes).")

        # 2. Validate filename and extension
        _, ext = os.path.splitext(filename.lower())
        if ext not in cls.ALLOWED_EXTENSIONS:
            raise FileValidationError(
                f"Unsupported file type '{ext}'. Only {', '.join(cls.ALLOWED_EXTENSIONS)} files are supported."
            )

        # 3. Reject executable magic headers
        for magic, desc in cls.DANGEROUS_MAGIC_HEADERS:
            if file_bytes.startswith(magic):
                raise FileValidationError(f"Malicious file rejected: Detected {desc} header.")

        # 4. Safe text extraction based on file type
        if ext == ".txt":
            return cls._extract_txt(file_bytes, filename)
        elif ext == ".pdf":
            return cls._extract_pdf(file_bytes, filename)
        elif ext == ".docx":
            return cls._extract_docx(file_bytes, filename)
        else:
            raise FileValidationError(f"Unhandled file extension: {ext}")

    @classmethod
    def _extract_txt(cls, file_bytes: bytes, filename: str) -> ExtractedDocument:
        # Decode as UTF-8 or Latin-1 with replacement
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="replace")

        # Sanity check: ensure file is not binary
        non_printable_count = sum(1 for c in text[:1000] if ord(c) < 32 and c not in "\n\r\t")
        if non_printable_count > 50:
            raise FileValidationError("File appears to contain binary data rather than plain text.")

        return cls._create_doc(filename, "txt", len(file_bytes), text, page_count=1)

    @classmethod
    def _extract_pdf(cls, file_bytes: bytes, filename: str) -> ExtractedDocument:
        if not file_bytes.startswith(b"%PDF-"):
            raise FileValidationError("Invalid PDF: Missing %PDF- magic header.")

        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            max_pages = 50

            total_pages = len(reader.pages)
            for idx, page in enumerate(reader.pages[:max_pages]):
                page_content = page.extract_text() or ""
                pages_text.append(page_content)

            combined_text = "\n\n".join(pages_text).strip()
            return cls._create_doc(
                filename,
                "pdf",
                len(file_bytes),
                combined_text,
                page_count=min(total_pages, max_pages),
                metadata={"total_pages": total_pages, "scanned_pages": min(total_pages, max_pages)}
            )
        except Exception as e:
            raise FileValidationError(f"Failed to extract PDF text safely: {str(e)}")

    @classmethod
    def _extract_docx(cls, file_bytes: bytes, filename: str) -> ExtractedDocument:
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise FileValidationError("Invalid DOCX: Missing ZIP/DOCX magic header.")

        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

            # Also extract text inside tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            combined_text = "\n\n".join(paragraphs).strip()
            return cls._create_doc(
                filename,
                "docx",
                len(file_bytes),
                combined_text,
                page_count=1,
                metadata={"paragraphs_count": len(doc.paragraphs), "tables_count": len(doc.tables)}
            )
        except Exception as e:
            raise FileValidationError(f"Failed to extract DOCX text safely: {str(e)}")

    @classmethod
    def _create_doc(
        cls,
        filename: str,
        file_type: str,
        size_bytes: int,
        text: str,
        page_count: int,
        metadata: Optional[Dict[str, Any]] = None
    ) -> ExtractedDocument:
        words = re.findall(r"\w+", text)
        excerpt = text[:300].strip() + ("..." if len(text) > 300 else "")
        return ExtractedDocument(
            filename=filename,
            file_type=file_type,
            size_bytes=size_bytes,
            text=text,
            page_count=page_count,
            word_count=len(words),
            character_count=len(text),
            excerpt=excerpt,
            metadata=metadata or {}
        )
